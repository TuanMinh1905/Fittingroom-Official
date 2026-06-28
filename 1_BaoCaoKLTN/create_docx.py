import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Setup styles for Normal
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(13)

# Setup margins
for section in doc.sections:
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)

def add_heading(text, level=1):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.bold = True
    r.font.size = Pt(14)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def add_item(title):
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.font.bold = True
    r.font.size = Pt(13)
    p.paragraph_format.line_spacing = 1.5
    return p

def add_content(label, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    
    r_label = p.add_run(label + ":\n")
    # r_label.font.bold = True # removed bold for label to match draft
    r_label.font.italic = True
    
    r_text = p.add_run(text)
    return p

add_heading("CHƯƠNG 2: CƠ SỞ LÝ THUYẾT", 1)
doc.add_paragraph()

add_heading("2.1. TỔNG QUAN VỀ XÂY DỰNG GIAO DIỆN", 2)

add_item("2.1.1. Ngôn ngữ lập trình TypeScript")
add_content("Lý thuyết", "TypeScript mở rộng JavaScript với hệ thống kiểu tĩnh, hỗ trợ kiểm tra lỗi ngay khi viết code, đồng thời vẫn transpile về JavaScript để chạy trên trình duyệt.")
add_content("Vai trò trong Project", "Khai báo kiểu cho các thực thể hệ thống 3D và các API request/response. Giúp giảm thiểu lỗi trong quá trình trao đổi dữ liệu giữa frontend và backend, đặc biệt với các dữ liệu phức tạp của 3D mesh và mô hình TailorNet.")

add_item("2.1.2. Next.js (App Router)")
add_content("Lý thuyết", "Next.js là framework React full-stack, cung cấp mô hình Server-Side Rendering (SSR) và Client-Side Rendering linh hoạt, tối ưu hóa hiệu suất và SEO.")
add_content("Vai trò trong project", "Tổ chức giao diện phòng thử đồ ảo (Virtual Fitting Room) và các trang thương mại điện tử. Quản lý hệ thống routing cho người dùng, giúp tối ưu hóa việc tải các tài nguyên 3D lớn và giữ cho giao diện mượt mà.")

add_item("2.1.3. Zustand (State Management)")
add_content("Lý thuyết", "Zustand là một thư viện quản lý state nhỏ gọn, nhanh và linh hoạt cho React. Nó dựa trên hook, không cần bọc ứng dụng trong Provider như Context hay Redux.")
add_content("Vai trò trong Project", "Quản lý trạng thái toàn cục cho phòng thử đồ (ví dụ: mô hình đang chọn, giới tính avatar, trạng thái loading của mesh, camera). Giúp các component 3D (Three.js) và UI (React) đồng bộ dữ liệu dễ dàng.")

add_item("2.1.4. Tailwind CSS")
add_content("Lý thuyết", "Tailwind CSS là CSS framework dạng utility-first, cung cấp các class nhỏ để xây dựng giao diện nhanh chóng.")
add_content("Vai trò trong project", "Xây dựng giao diện bảng điều khiển, các nút chọn trang phục, slider điều chỉnh camera trong phòng thử đồ 3D. Đảm bảo giao diện responsive trên mọi thiết bị và đồng nhất.")

add_item("2.1.5. Three.js và React Three Fiber")
add_content("Lý thuyết", "Three.js là thư viện đồ họa 3D mạnh mẽ chạy trên trình duyệt sử dụng WebGL. React Three Fiber là một renderer cho Three.js trong môi trường React.")
add_content("Vai trò trong project", "Dùng để render mô hình cơ thể người (Avatar) và các trang phục (áo, quần) dạng lưới 3D (Mesh) ngay trên trình duyệt. Xử lý các ánh sáng, chất liệu (materials), và camera để người dùng có thể quan sát từ nhiều góc độ.")

add_heading("2.2. TỔNG QUAN VỀ XÂY DỰNG CHỨC NĂNG VÀ MÔ HÌNH HÓA 3D", 2)

add_item("2.2.1. HonoJS (Backend API)")
add_content("Lý thuyết", "Hono là một web framework nhỏ, cực kỳ nhanh, chạy được trên nhiều runtime (Cloudflare Workers, Deno, Bun, Node.js).")
add_content("Vai trò trong project", "Xử lý các API backend cho ứng dụng, điều phối luồng dữ liệu giữa frontend (Next.js) và các dịch vụ học máy (Python). Quản lý logic nghiệp vụ, xác thực và phân phối tài nguyên.")

add_item("2.2.2. FastAPI (Xử lý học máy Python)")
add_content("Lý thuyết", "FastAPI là một web framework hiện đại, hiệu năng cao để xây dựng các API bằng Python.")
add_content("Vai trò trong project", "Xây dựng các microservice chuyên biệt cho xử lý mô hình AI. Cung cấp API để nhận các thông số trang phục và sinh ra lưới 3D, xử lý giao tiếp đa luồng để tránh xung đột tài nguyên khi suy luận mô hình.")

add_item("2.2.3. SMPL (Skinned Multi-Person Linear Model)")
add_content("Lý thuyết", "SMPL là mô hình toán học 3D biểu diễn cơ thể người với khả năng thay đổi hình dáng (shape) và tư thế (pose) thông qua các tham số linh hoạt.")
add_content("Vai trò trong project", "Cung cấp khung cơ thể (Avatar) nền tảng trong không gian 3D. Các thông số SMPL được dùng để điều chỉnh kích thước cơ thể người dùng (nam/nữ, béo/gầy, cao/thấp) để mặc thử đồ ảo chính xác.")

add_item("2.2.4. TailorNet")
add_content("Lý thuyết", "TailorNet là mô hình học sâu có khả năng dự đoán biến dạng của quần áo 3D dựa trên tư thế, hình dáng cơ thể và kiểu dáng trang phục. Nó xử lý cả các chi tiết nếp nhăn thực tế.")
add_content("Vai trò trong project", "Lõi công nghệ AI của hệ thống. Nhận dữ liệu thông số cơ thể từ SMPL để sinh ra lưới 3D (mesh) của áo và quần. Giúp trang phục ôm sát vào cơ thể ảo một cách chân thực nhất, giải quyết bài toán chống xuyên thấu giữa các lớp.")

output_file = 'Chuong2_CoSoLyThuyet_Project.docx'
doc.save(output_file)
print(f"Saved to {output_file}")
