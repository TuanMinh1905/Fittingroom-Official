import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Default style
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(13)
# Paragraph spacing
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# Document margins
for section in doc.sections:
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)

def add_heading_1(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.bold = True
    r.font.size = Pt(14)
    p.paragraph_format.space_after = Pt(12)
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.bold = True
    r.font.size = Pt(13)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_heading_3(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.bold = True
    r.font.size = Pt(13)
    p.paragraph_format.space_before = Pt(6)
    return p

def add_paragraph_with_title(title, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(1.27)
    
    r_title = p.add_run(title + " ")
    r_title.font.italic = True
    
    r_text = p.add_run(text)
    return p

def add_normal_paragraph(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Cm(1.27)
    return p

# CHAPTER 2
add_heading_1("CHƯƠNG 2: CƠ SỞ LÝ THUYẾT")

# 2.1
add_heading_2("2.1. TỔNG QUAN VỀ XÂY DỰNG GIAO DIỆN")

add_heading_3("2.1.1. Ngôn ngữ lập trình TypeScript")
add_paragraph_with_title("Lý thuyết:", "TypeScript mở rộng JavaScript với hệ thống kiểu tĩnh, hỗ trợ kiểm tra lỗi ngay trong quá trình viết mã (compile-time), đồng thời mã nguồn vẫn được biên dịch (transpile) về JavaScript để chạy trên mọi trình duyệt hoặc môi trường Node.js. TypeScript giúp tăng tính chặt chẽ, dễ bảo trì và mở rộng cho các dự án quy mô lớn.")
add_paragraph_with_title("Vai trò trong Project:", "Trong dự án này, TypeScript được sử dụng để khai báo kiểu tĩnh cho các thực thể 3D, hệ thống camera, các API request và response. Việc này giúp giảm thiểu tối đa lỗi trong quá trình trao đổi dữ liệu phức tạp giữa frontend và backend, đặc biệt là khi làm việc với các mảng dữ liệu điểm (vertices) của 3D mesh và thông số từ mô hình TailorNet.")

add_heading_3("2.1.2. Next.js (App Router)")
add_paragraph_with_title("Lý thuyết:", "Next.js là một framework React full-stack phổ biến, cung cấp mô hình Server-Side Rendering (SSR), Static Site Generation (SSG) và Client-Side Rendering linh hoạt. Với kiến trúc App Router mới, Next.js tối ưu hóa quá trình tải trang, hỗ trợ routing tự động theo thư mục và tối ưu hóa SEO mạnh mẽ.")
add_paragraph_with_title("Vai trò trong Project:", "Next.js đóng vai trò tổ chức toàn bộ giao diện của phòng thử đồ ảo (Virtual Fitting Room). Nó quản lý hệ thống định tuyến (routing) cho người dùng, tối ưu hóa việc tải trước (prefetch) các tài nguyên 3D lớn (như file .obj, .gltf) và giữ cho giao diện tương tác luôn mượt mà trong quá trình người dùng thay đổi trang phục hay giới tính avatar.")

add_heading_3("2.1.3. Zustand (State Management)")
add_paragraph_with_title("Lý thuyết:", "Zustand là một thư viện quản lý trạng thái (state management) nhỏ gọn, nhanh và linh hoạt dành cho React. Nó hoạt động dựa trên các hook mà không cần bọc ứng dụng trong Provider như Context API hay Redux, giúp giảm thiểu boilerplate code và tăng hiệu năng.")
add_paragraph_with_title("Vai trò trong Project:", "Zustand được sử dụng để quản lý trạng thái toàn cục của hệ thống phòng thử đồ ảo. Các trạng thái này bao gồm: mô hình 3D đang được chọn, giới tính của avatar (nam/nữ), trạng thái tải (loading) của lưới (mesh), và các thông số cấu hình camera. Nhờ Zustand, các component 3D (Three.js) và các component giao diện UI (React) có thể đồng bộ dữ liệu với nhau một cách dễ dàng và hiệu quả.")

add_heading_3("2.1.4. Tailwind CSS")
add_paragraph_with_title("Lý thuyết:", "Tailwind CSS là một CSS framework theo hướng utility-first, cung cấp sẵn hàng loạt các class tiện ích nhỏ (như margin, padding, color, flexbox, grid, v.v.) để xây dựng trực tiếp giao diện trên file HTML/JSX mà không cần phải viết CSS thuần phức tạp.")
add_paragraph_with_title("Vai trò trong Project:", "Tailwind CSS được dùng để thiết kế toàn bộ giao diện bảng điều khiển, các nút chọn trang phục, thanh trượt (slider) điều chỉnh camera trong phòng thử đồ 3D. Nhờ Tailwind CSS, giao diện của ứng dụng được đảm bảo tính responsive trên nhiều kích thước màn hình thiết bị khác nhau, đồng thời giữ được sự thống nhất về phong cách thiết kế trên toàn bộ dự án.")

add_heading_3("2.1.5. Three.js và React Three Fiber")
add_paragraph_with_title("Lý thuyết:", "Three.js là một thư viện đồ họa 3D mạnh mẽ chạy trên trình duyệt web thông qua WebGL. React Three Fiber (R3F) là một renderer dành cho Three.js hoạt động trong môi trường React, cho phép khai báo các đối tượng 3D dưới dạng các component React tái sử dụng được.")
add_paragraph_with_title("Vai trò trong Project:", "Bộ đôi này là thành phần cốt lõi để hiển thị môi trường 3D. Chúng được dùng để render mô hình cơ thể người (Avatar) và các trang phục (áo, quần) dưới dạng lưới 3D (Mesh) ngay trên trình duyệt web. R3F kết hợp Three.js cũng chịu trách nhiệm xử lý các nguồn sáng (lighting), chất liệu bề mặt (materials) và hệ thống camera, cho phép người dùng quan sát trang phục từ nhiều góc độ khác nhau một cách chân thực nhất.")

# 2.2
add_heading_2("2.2. TỔNG QUAN VỀ XÂY DỰNG CHỨC NĂNG VÀ MÔ HÌNH HÓA 3D")

add_heading_3("2.2.1. HonoJS (Backend API)")
add_paragraph_with_title("Lý thuyết:", "Hono là một web framework siêu nhỏ và cực kỳ nhanh, có khả năng chạy trên nhiều môi trường runtime đa dạng như Cloudflare Workers, Deno, Bun và Node.js. Hono nổi bật với hiệu suất cao và cú pháp hiện đại, tối giản.")
add_paragraph_with_title("Vai trò trong Project:", "HonoJS xử lý các API backend chính cho ứng dụng, đóng vai trò là một API Gateway điều phối luồng dữ liệu giữa frontend (Next.js) và các dịch vụ học máy viết bằng Python. Nó phụ trách quản lý logic nghiệp vụ, xử lý xác thực (authentication), quản lý phiên bản trang phục và phân phối tài nguyên 3D cho client một cách nhanh chóng.")

add_heading_3("2.2.2. FastAPI (Xử lý học máy Python)")
add_paragraph_with_title("Lý thuyết:", "FastAPI là một web framework hiện đại, hiệu năng rất cao dùng để xây dựng các API bằng Python, dựa trên tiêu chuẩn gợi ý kiểu (type hints) của Python. Nhờ tích hợp Pydantic và Starlette, FastAPI cung cấp tốc độ phản hồi nhanh chóng ngang ngửa với NodeJS và Go.")
add_paragraph_with_title("Vai trò trong Project:", "FastAPI được sử dụng để xây dựng các microservice chuyên biệt cho việc xử lý suy luận (inference) mô hình AI. Cụ thể, nó cung cấp các API để nhận thông số trang phục từ người dùng, đưa qua các mạng nơ-ron để sinh ra lưới 3D (3D mesh). Ngoài ra, FastAPI còn xử lý cơ chế giao tiếp đa luồng (multi-threading) và quản lý bộ nhớ đệm (LRU cache) cho các luồng xử lý TailorNet, giúp tránh xung đột tài nguyên GPU khi nhiều request đến cùng lúc.")

add_heading_3("2.2.3. SMPL (Skinned Multi-Person Linear Model)")
add_paragraph_with_title("Lý thuyết:", "SMPL là một mô hình toán học biểu diễn cơ thể người 3D hiện đại, có khả năng thay đổi hình dáng (shape) và tư thế (pose) thông qua tập hợp các tham số tuyến tính. SMPL được huấn luyện dựa trên hàng ngàn bản quét 3D cơ thể người thực tế để đưa ra mô phỏng biến dạng cơ bắp và da chính xác nhất.")
add_paragraph_with_title("Vai trò trong Project:", "SMPL cung cấp khung cơ thể (Avatar) nền tảng trong không gian 3D của ứng dụng. Các thông số của SMPL (shape, pose parameters) được hệ thống sử dụng để điều chỉnh kích thước cơ thể người dùng ảo (như phân biệt nam/nữ, tỉ lệ béo/gầy, chiều cao), qua đó tạo ra một hình nhân ảo có thông số vật lý cá nhân hóa để mặc thử đồ chính xác.")

add_heading_3("2.2.4. TailorNet")
add_paragraph_with_title("Lý thuyết:", "TailorNet là một mô hình học sâu chuyên biệt có khả năng dự đoán biến dạng của quần áo 3D dựa trên sự thay đổi của tư thế (pose), hình dáng cơ thể (shape) và kiểu dáng trang phục (style). Mô hình này xử lý xuất sắc các chi tiết nhỏ như nếp nhăn thực tế trên vải khi cơ thể chuyển động, mang lại kết quả mô phỏng vượt trội so với các phương pháp vật lý truyền thống.")
add_paragraph_with_title("Vai trò trong Project:", "TailorNet chính là lõi công nghệ trí tuệ nhân tạo của hệ thống Fitting Room. Mô hình nhận dữ liệu thông số cơ thể từ SMPL để sinh ra lưới 3D (mesh) của áo và quần. Nhờ khả năng dự đoán nếp nhăn và biến dạng bề mặt, TailorNet giúp trang phục ôm sát vào cơ thể ảo một cách chân thực nhất, đồng thời tích hợp các thuật toán giảm thiểu lỗi xuyên thấu (interpenetration) giữa bề mặt trang phục và cơ thể hoặc giữa các lớp trang phục (ví dụ: áo và quần).")

doc.save('Chuong2_CoSoLyThuyet_Project.docx')
