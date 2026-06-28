"""
Script tạo file KLTN Chương 2 - Cơ Sở Lý Thuyết
Nội dung CỤ THỂ theo source code của project TMF-NodeJS-FittingRoom3D
"""
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# --- Default Style ---
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(13)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# --- Margins ---
for section in doc.sections:
    section.top_margin    = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin   = Cm(3.0)
    section.right_margin  = Cm(2.0)

INDENT = Cm(1.27)

def h1(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.bold = True
    r.font.size = Pt(14)
    p.paragraph_format.space_after = Pt(10)

def h2(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.bold = True
    r.font.size = Pt(13)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(4)

def h3(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.bold = True
    r.font.size = Pt(13)
    p.paragraph_format.space_before = Pt(6)

def body(label, text):
    """Đoạn văn có nhãn in nghiêng + nội dung, thụt đầu dòng."""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = INDENT
    p.paragraph_format.line_spacing = 1.5
    rl = p.add_run(label + ": ")
    rl.font.italic = True
    p.add_run(text)

def body_plain(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = INDENT
    p.paragraph_format.line_spacing = 1.5

# =====================================================================
# CHƯƠNG 2
# =====================================================================
h1("CHƯƠNG 2: CƠ SỞ LÝ THUYẾT")

# -------------------------------------------------------------------
h2("2.1. TỔNG QUAN VỀ XÂY DỰNG GIAO DIỆN")

# 2.1.1 TypeScript
h3("2.1.1. Ngôn ngữ lập trình TypeScript")
body("Lý thuyết",
     "TypeScript mở rộng JavaScript bằng hệ thống kiểu tĩnh (static typing). "
     "Mã nguồn TypeScript được trình biên dịch (tsc) chuyển đổi (transpile) về JavaScript "
     "thuần trước khi chạy trên trình duyệt hoặc Node.js. "
     "TypeScript giúp phát hiện lỗi ngay tại thời điểm viết mã, tăng khả năng tái sử dụng "
     "và bảo trì cho dự án quy mô lớn.")
body("Vai trò trong Project",
     "Toàn bộ frontend (frontend-next/) và backend (backend-hono/) của dự án đều viết bằng TypeScript. "
     "Trong file frontend-next/src/store/productStore.ts, TypeScript được dùng để định nghĩa kiểu "
     "Product (bao gồm _id, name, slug, price, sizeChart, garment_type,...) và SizeMeasurement. "
     "Điều này đảm bảo khi fetchProducts() gọi API http://localhost:8003/products, "
     "dữ liệu trả về được ép kiểu Product[] ngay lập tức, giúp tránh lỗi truy cập sai tên field "
     "trong toàn bộ ứng dụng. "
     "Trong backend-hono/src/models/user.ts, interface IUser kế thừa Document của Mongoose, "
     "định nghĩa chặt chẽ các trường name, email, password, role, address, phone, gender, birthday. "
     "Trong backend-hono/src/models/product.ts, interface ISizeMeasurement khai báo kiểu cho "
     "từng dòng bảng số đo (size, length_cm, chest_half_cm, shoulder_cm, waist_cm, hip_cm) "
     "để backend-hono và frontend đồng bộ cùng một cấu trúc dữ liệu.")

# 2.1.2 Next.js
h3("2.1.2. Next.js (App Router)")
body("Lý thuyết",
     "Next.js là framework React full-stack hỗ trợ Server-Side Rendering (SSR), "
     "Static Site Generation (SSG) và Client-Side Rendering linh hoạt. "
     "Với kiến trúc App Router (thư mục app/), mỗi thư mục con tương ứng với một route URL, "
     "Next.js tự động code-splitting và prefetch tài nguyên theo route.")
body("Vai trò trong Project",
     "Frontend của dự án được tổ chức trong thư mục frontend-next/src/app/(public)/. "
     "Các route chính gồm: "
     "app/(public)/page.tsx là trang chủ; "
     "app/(public)/fitting-room/ là phòng thử đồ ảo 3D; "
     "app/(public)/p/ là trang chi tiết sản phẩm; "
     "app/(public)/cart/ là giỏ hàng; "
     "app/(public)/checkout/ là đặt hàng; "
     "app/(public)/brand/, category/, product/, blog/ là các trang liệt kê. "
     "Trong component FittingRoom3DViewer.tsx (frontend-next/src/component/fittingroom/), "
     "Next.js dynamic import với tùy chọn ssr: false được dùng để lazy-load Scene 3D "
     "(import từ sectionSMPL/Sence.tsx), vì Three.js cần đối tượng window và WebGL "
     "chỉ tồn tại phía trình duyệt, không có phía server. "
     "Điều này giúp trang tải nhanh hơn bằng cách trì hoãn việc tải thư viện 3D nặng (~1MB).")

# 2.1.3 Zustand
h3("2.1.3. Zustand (State Management)")
body("Lý thuyết",
     "Zustand là thư viện quản lý trạng thái nhỏ gọn cho React, hoạt động dựa trên hook "
     "mà không cần bọc ứng dụng trong Provider như Redux hay Context API. "
     "Mỗi store được tạo bằng hàm create<T>() và có thể được subscribe từ bất kỳ component nào.")
body("Vai trò trong Project",
     "Dự án có 6 store trong frontend-next/src/store/: "
     "(1) productStore.ts: quản lý danh sách Product[], hàm fetchProducts() gọi "
     "http://localhost:8003/products, fetchProductsByCategory(categorySlug), "
     "fetchProductsByBrand(brandSlug); "
     "(2) cartStore.ts: quản lý giỏ hàng với persist middleware của Zustand, "
     "lưu giỏ hàng vào localStorage dưới key 'tmf-cart-storage', "
     "cung cấp addToCart(), removeFromCart(), updateQuantity(), "
     "getTotalItems(), getTotalPrice() cho toàn bộ ứng dụng; "
     "(3) fittingRoom.ts: quản lý dữ liệu mesh 3D trả về từ "
     "backend-smpl (http://localhost:8001/smpl), lưu fittingRoomData gồm "
     "vertices[], faces[], n_vertices, n_faces; "
     "(4) categoryStore.ts, brandStore.ts, blogStore.ts: quản lý dữ liệu "
     "tương ứng cho từng tính năng. "
     "Nhờ Zustand, component FittingRoom3DViewer.tsx và Scene.tsx (Three.js) "
     "có thể đọc cùng một nguồn mesh data mà không cần truyền props qua nhiều tầng.")

# 2.1.4 Tailwind CSS
h3("2.1.4. Tailwind CSS")
body("Lý thuyết",
     "Tailwind CSS là CSS framework theo hướng utility-first, cung cấp sẵn các class tiện ích "
     "(margin, padding, color, flex, grid,...) để xây dựng giao diện trực tiếp trong JSX "
     "mà không cần viết CSS thuần. Tailwind v4 trong dự án này được tích hợp qua "
     "@tailwindcss/postcss và cấu hình trong postcss.config.mjs.")
body("Vai trò trong Project",
     "Tailwind CSS được dùng để xây dựng toàn bộ giao diện trong frontend-next/src/. "
     "Trong FittingRoom3DViewer.tsx, Tailwind class được dùng để bố cục khung 3D viewer "
     "(w-full h-full relative flex), trạng thái loading (animate-spin, text-blue-400), "
     "trạng thái lỗi (text-red-400, bg-red-900/30, border-red-800/50). "
     "Component HumanMesh.tsx và GarmentMesh.tsx dùng Tailwind cho wrapper div. "
     "Thanh trượt dọc (vertical slider) trong FittingRoom3DViewer.tsx có style CSS tùy chỉnh "
     "kết hợp với Tailwind để tạo hiệu ứng gradient và animation glow cho slider thumb.")

# 2.1.5 Three.js và React Three Fiber
h3("2.1.5. Three.js và React Three Fiber (R3F)")
body("Lý thuyết",
     "Three.js là thư viện đồ họa 3D chạy trên trình duyệt thông qua WebGL. "
     "React Three Fiber (R3F) là renderer cho Three.js trong môi trường React, "
     "cho phép khai báo đối tượng 3D dưới dạng JSX component. "
     "@react-three/drei là tập hợp các helper component sẵn có như OrbitControls, ContactShadows. "
     "Dự án dùng phiên bản three@0.183.2, @react-three/fiber@9.5.0, @react-three/drei@10.7.7.")
body("Vai trò trong Project",
     "Trong frontend-next/src/component/fittingroom/sectionSMPL/Sence.tsx, "
     "component <Canvas> của R3F tạo WebGL context với camera position [0, 0.9, 2.5], fov 40 độ "
     "và background gradient tối. Hệ thống đèn 3-point được thiết lập: "
     "ambientLight (intensity 0.4), directionalLight key light (position [3,4,3], intensity 1.5), "
     "fill light ([-3,2,1], intensity 0.6) và rim light ([0,3,-3], intensity 0.4). "
     "Component HumanMesh.tsx (frontend-next/src/component/fittingroom/sectionSMPL/HumanMesh.tsx) "
     "nhận mảng vertices[] và faces[] phẳng 1 chiều từ API, dùng THREE.BufferGeometry() "
     "với BufferAttribute để dựng mesh cơ thể người (SMPL) với màu da #e8beac. "
     "Component GarmentMesh.tsx nhận thêm tham số type (t-shirt/pant/...) "
     "để áp dụng Stencil Buffer: áo (upper) render trước với stencilFunc=AlwaysStencilFunc, "
     "ghi stencil=1 vào buffer; quần (lower) chỉ render ở pixel stencil≠1 "
     "(NotEqualStencilFunc), tránh hiện tượng mesh xuyên nhau tại vùng eo/hông. "
     "SmartOrbitControls và CameraTargetController cho phép người dùng xoay, zoom, "
     "và điều khiển camera target theo thanh trượt dọc (min Y=-1.0, max Y=1.6).")

# -------------------------------------------------------------------
h2("2.2. TỔNG QUAN VỀ XÂY DỰNG CHỨC NĂNG VÀ MÔ HÌNH HÓA 3D")

# 2.2.1 HonoJS
h3("2.2.1. HonoJS (Backend API)")
body("Lý thuyết",
     "Hono là web framework siêu nhỏ gọn và hiệu năng cao, chạy trên nhiều runtime "
     "như Node.js, Bun, Deno, Cloudflare Workers. "
     "Hono cung cấp cú pháp routing tương tự Express nhưng nhanh hơn đáng kể "
     "và hỗ trợ TypeScript native.")
body("Vai trò trong Project",
     "HonoJS là backend API chính, chạy tại http://localhost:8003 (cấu hình trong "
     "backend-hono/src/index.ts, biến PORT). "
     "File index.ts khởi tạo app Hono, bật CORS cho origin http://localhost:3000, "
     "kết nối MongoDB qua hàm connectDB() (backend-hono/src/db/mongoose.ts), "
     "rồi gắn các router con: "
     "app.route('/products', products) – CRUD sản phẩm, tìm kiếm theo tên (/products/search?q=), "
     "lọc theo categorySlug (/products/category/:categorySlug) và brandSlug (/products/brand/:brandSlug); "
     "app.route('/category', category) – quản lý danh mục; "
     "app.route('/brand', brand) – quản lý thương hiệu; "
     "app.route('/blog', blog) – quản lý bài viết; "
     "app.route('/auth', auth) – đăng nhập và đăng ký; "
     "app.route('/users', user) – quản lý người dùng; "
     "app.route('/orders', order) – tạo đơn hàng, cập nhật trạng thái, thống kê doanh thu; "
     "app.route('/ai-advisor', aiAdvisor) – proxy gọi OpenAI GPT-4o-mini tư vấn size trang phục. "
     "Trong backend-hono/src/routes/ai-advisor.ts, route POST /ai-advisor nhận "
     "message + context + history từ frontend, gọi openai.chat.completions.create() "
     "với model gpt-4o-mini, system prompt chứa bảng size chart TMF (S/M/L/XL cho áo và quần, "
     "cả nam lẫn nữ) và trả về lời tư vấn bằng tiếng Việt.")

# 2.2.2 FastAPI + MongoDB
h3("2.2.2. FastAPI và MongoDB (Mongoose)")
body("Lý thuyết FastAPI",
     "FastAPI là web framework Python hiệu năng cao, dựa trên Starlette và Pydantic. "
     "FastAPI tự động sinh OpenAPI docs (/docs), validate request body qua Pydantic BaseModel "
     "và hỗ trợ async/await natively.")
body("Vai trò FastAPI trong Project",
     "Dự án có hai FastAPI service riêng biệt: "
     "(1) backend-smpl/ (chạy tại cổng 8001): file main.py và smpl_model.py "
     "cung cấp endpoint POST /smpl nhận betas: list[float] từ frontend, "
     "tính toán v_shaped = v_template + shapedirs @ betas (công thức SMPL), "
     "trả về vertices[], faces[], n_vertices=6890, n_faces=13776 cho Three.js render. "
     "Khi khởi động, server dùng @asynccontextmanager lifespan để pre-load "
     "file smpl_male.pkl vào bộ nhớ (backend-smpl/models/); "
     "(2) TailorNet/TailorNet-master/backend/main.py (chạy trong Docker): "
     "cung cấp POST /try-on nhận số đo chiều cao, cân nặng, giới tính, garment_type, size "
     "và POST /try-on-outfit để thử nhiều trang phục cùng lúc; "
     "POST /switch-gender để chuyển đổi giới tính avatar và evict model cũ khỏi RAM; "
     "GET /health kiểm tra trạng thái cached models.")
body("Lý thuyết MongoDB + Mongoose",
     "MongoDB là hệ quản trị cơ sở dữ liệu NoSQL document-based. "
     "Mongoose là ODM (Object Data Modeling) cho MongoDB trong Node.js, "
     "cho phép định nghĩa Schema, tạo Model và thực hiện CRUD với cú pháp TypeScript.")
body("Vai trò MongoDB + Mongoose trong Project",
     "File backend-hono/src/db/mongoose.ts gọi mongoose.connect() kết nối đến "
     "mongodb://localhost:27017/TMFDatabaseOfficial (cấu hình qua MONGODB_URL trong .env). "
     "Dự án định nghĩa 6 Schema trong backend-hono/src/models/: "
     "ProductSchema (name, slug, price, discountPrice, imageUrl, categorySlug, rating, "
     "soldCount, brand, stock, sizes[], colorCodes[], sizeChart[], garment_type); "
     "UserSchema (name, email, password select:false, role: user|admin, address, phone, gender, birthday) "
     "với middleware pre('save') dùng bcryptjs.genSalt(10) + bcryptjs.hash() "
     "để mã hóa mật khẩu tự động trước khi lưu; "
     "OrderSchema với các trường items[], totalPrice, shippingFee, paymentMethod, "
     "status: Pending|Processing|Delivered|Cancelled; "
     "CategorySchema, BrandSchema, BlogSchema. "
     "Route POST /auth/login (backend-hono/src/routes/auth.ts) nhận email + password, "
     "gọi User.findOne({ email }).select('+password') để lấy hash, "
     "rồi gọi user.comparePassword(password) – hàm này dùng bcrypt.compare() – "
     "nếu đúng trả về thông tin user (đã xóa trường password). "
     "Route GET /orders/stats (backend-hono/src/routes/order.ts) tổng hợp "
     "totalRevenue, statusBreakdown, monthlyRevenue (12 tháng gần nhất), "
     "bestSellers và uniqueCustomers cho dashboard Admin.")

# 2.2.3 SMPL
h3("2.2.3. SMPL (Skinned Multi-Person Linear Model)")
body("Lý thuyết",
     "SMPL là mô hình toán học biểu diễn cơ thể người 3D, được xây dựng từ hàng nghìn bản quét "
     "3D thực tế. SMPL có hai tập tham số: beta (shape parameters, 10 chiều) điều chỉnh hình dáng "
     "cơ thể và theta (pose parameters, 72 chiều) điều chỉnh tư thế. "
     "Công thức tính mesh là: v_shaped = v_template + Σ(shapedirs_i × beta_i), "
     "trong đó v_template là mesh cơ thể trung bình (6890 đỉnh, 13776 mặt tam giác).")
body("Vai trò trong Project",
     "SMPL được triển khai trong hai tầng của dự án: "
     "(1) Tầng cơ bản – backend-smpl/smpl_model.py: class SMPLModel load file "
     "smpl_male.pkl bằng pickle.load(), xử lý tương thích chumpy (thư viện cũ, không chạy "
     "trên Python >= 3.12) bằng cách giả lập chumpy.Ch. "
     "Hàm get_mesh(betas) tính v_shaped qua np.einsum('ijk,k->ij', shapedirs, betas), "
     "flatten thành mảng 1D và trả về cho frontend Three.js. "
     "Hàm forward() giới hạn giá trị betas trong [-3.0, 3.0] bằng np.clip() "
     "và đảm bảo đúng 10 phần tử bằng np.pad(); "
     "(2) Tầng nâng cao – TailorNet/TailorNet-master/: dùng SMPL4Garment "
     "(models/smpl4garment.py) để tính body mesh và garment template mesh "
     "từ beta + theta + garment_d (vector biến dạng từ TailorNet). "
     "Hàm get_smpl(gender) trong backend/main.py cache SMPL instance riêng cho "
     "male và female (file model.pkl lưu tại /app/data/smpl_models/{gender}/model.pkl trong Docker), "
     "tự động fallback sang neutral nếu file không tồn tại.")

# 2.2.4 TailorNet
h3("2.2.4. TailorNet")
body("Lý thuyết",
     "TailorNet là mô hình học sâu dự đoán biến dạng 3D của quần áo dựa trên ba tham số: "
     "theta (tư thế cơ thể), beta (hình dáng cơ thể) và gamma (kiểu dáng trang phục/style). "
     "Mô hình được huấn luyện riêng cho từng loại trang phục (t-shirt, shirt, pant, short-pant, skirt) "
     "và từng giới tính. Kiến trúc gồm ba mạng: SS2G (shape-to-garment), "
     "LF (low-frequency deformation) và HF (high-frequency wrinkles).")
body("Vai trò trong Project",
     "TailorNet là lõi AI của hệ thống Fitting Room, được triển khai trong "
     "TailorNet/TailorNet-master/. "
     "File backend/main.py quản lý TailorNet runner qua hàm get_tailornet_runner(garment_class, gender), "
     "áp dụng chiến lược Exclusive Resource: chỉ phục vụ 1 gender tại 1 thời điểm. "
     "Khi người dùng chuyển giới tính, endpoint POST /switch-gender gọi _evict_all_runners() "
     "để xóa TẤT CẢ model cũ khỏi RAM, sau đó gọi _force_free_memory() "
     "(malloc_trim trên Linux) và pre-load SMPL mới. "
     "Cơ chế threading.Lock() (_model_lock) đảm bảo chỉ 1 luồng load model tại một thời điểm "
     "tránh race condition và Out-Of-Memory. "
     "Endpoint POST /try-on thực hiện quy trình: "
     "(1) measurements_to_beta() chuyển chiều cao + cân nặng → beta 10 chiều; "
     "(2) garment_measurements_to_gamma() chuyển số đo thực tế của sản phẩm (từ sizeChart trong MongoDB) "
     "→ gamma vector; "
     "(3) tn_runner.forward(thetas, betas, gammas) sinh ra garment_d – vector biến dạng; "
     "(4) smpl.run() tạo body mesh và garment mesh từ garment_d; "
     "(5) remove_interpenetration_fast() (utils/interpenetration.py) đẩy các đỉnh garment "
     "ra khỏi bề mặt body để loại bỏ hiện tượng mesh xuyên nhau; "
     "(6) resolve_garment_overlap() xử lý xuyên mesh giữa áo và quần "
     "theo thứ tự layer: pant/skirt (layer 0) → t-shirt/shirt (layer 1); "
     "(7) trả về mesh_data gồm body_vertices, body_faces và garments[] "
     "cho frontend FittingRoom3DViewer.tsx render bằng Three.js, "
     "cùng image_small và image_large (ảnh PNG base64 render sẵn bằng PyOpenGL).")

doc.save('Chuong2_CoSoLyThuyet_Project_v2.docx')
print("Done! Saved: Chuong2_CoSoLyThuyet_Project_v2.docx")
