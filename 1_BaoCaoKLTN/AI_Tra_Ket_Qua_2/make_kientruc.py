"""
Tạo đoạn văn mô tả Kiến Trúc Hệ Thống TMF Virtual Fitting Room 3D.
Format bám sát draft.docx:
  - Heading 3: "Kiến trúc hệ thống"
  - [Hình placeholder]
  - Các đoạn văn Normal, justify, first_line_indent
"""

import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

for sec in doc.sections:
    sec.top_margin    = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin   = Cm(3.0)
    sec.right_margin  = Cm(2.0)

normal = doc.styles['Normal']
normal.font.name        = 'Times New Roman'
normal.font.size        = Pt(13)
normal.paragraph_format.line_spacing = 1.5

FI  = Cm(0.8)   # first_line_indent
LI  = Cm(0.8)   # left_indent for body

def heading3(text):
    p = doc.add_paragraph(style='Heading 3')
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    return p

def body(text):
    p = doc.add_paragraph()
    p.paragraph_format.alignment         = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing      = 1.5
    p.paragraph_format.first_line_indent = FI
    p.paragraph_format.left_indent       = LI
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    return p

def caption(text):
    """Caption ảnh - căn giữa, italic"""
    p = doc.add_paragraph()
    p.paragraph_format.alignment    = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    r.font.name  = 'Times New Roman'
    r.font.size  = Pt(13)
    r.italic     = True
    return p

def blank():
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5

# ══════════════════════════════════════════════════════════════════════════════
# NỘI DUNG
# ══════════════════════════════════════════════════════════════════════════════

heading3('Kiến trúc hệ thống')

blank()   # placeholder cho hình (chèn ảnh thủ công sau)

caption('Hình X. Mô hình kiến trúc của hệ thống TMF Virtual Fitting Room 3D.')

body(
    'Kiến trúc hệ thống của ứng dụng được thiết kế theo mô hình Client – Server '
    'kết hợp Microservice, nhằm đảm bảo khả năng mở rộng, hiệu năng và tính linh '
    'hoạt trong quá trình vận hành. Hệ thống được phân chia thành ba tầng chính: '
    'tầng giao diện người dùng (Frontend), tầng xử lý nghiệp vụ (Backend API) và '
    'tầng trí tuệ nhân tạo 3D (AI Engine), các tầng giao tiếp với nhau thông qua '
    'RESTful API chuẩn HTTP.'
)

body(
    'Ở phía Frontend, hệ thống sử dụng Next.js (React) để xây dựng giao diện '
    'website theo mô hình App Router với Server-Side Rendering, cho phép người dùng '
    'thực hiện các chức năng như duyệt danh mục sản phẩm, nhập số đo cơ thể, thử đồ '
    'ảo 3D và quản lý đơn hàng. Giao diện được xây dựng với TailwindCSS và TypeScript, '
    'trạng thái toàn cục được quản lý bởi Zustand. Tính năng phòng thử đồ 3D được '
    'hiện thực bằng Three.js kết hợp React Three Fiber, cho phép render mesh cơ thể '
    'và trang phục trực tiếp trên trình duyệt thông qua WebGL mà không cần cài đặt '
    'phần mềm bổ sung.'
)

body(
    'Phía Backend đóng vai trò trung tâm là API Server được xây dựng bằng HonoJS '
    'trên nền tảng Node.js, chịu trách nhiệm xử lý toàn bộ các request từ client, '
    'quản lý xác thực người dùng bằng JWT, điều phối luồng dữ liệu giữa các thành '
    'phần trong hệ thống và giao tiếp với các dịch vụ bên ngoài. API Server cũng là '
    'thành phần trung gian kết nối với OpenAI API (GPT-4o-mini) để xử lý tác vụ tư '
    'vấn chọn size trang phục bằng tiếng Việt, dựa trên số đo cơ thể người dùng và '
    'bảng size của từng sản phẩm.'
)

body(
    'Hệ thống sử dụng MongoDB (thông qua Mongoose ODM) làm cơ sở dữ liệu chính để '
    'lưu trữ thông tin sản phẩm, người dùng, đơn hàng và bảng số đo sizeChart. '
    'MongoDB được lựa chọn vì tính linh hoạt của schema document-based, phù hợp với '
    'cấu trúc dữ liệu sản phẩm đa dạng của hệ thống thương mại điện tử thời trang. '
    'Dữ liệu được lưu trữ trên MongoDB Atlas Cloud với đầy đủ tính năng replica set '
    'đảm bảo tính sẵn sàng cao.'
)

body(
    'Đối với xử lý trí tuệ nhân tạo 3D, hệ thống sử dụng FastAPI (Python) được '
    'container hóa bằng Docker làm AI Engine riêng biệt. Khi nhận được số đo chiều '
    'cao, cân nặng và giới tính từ API Server, FastAPI sẽ chạy mô hình SMPL để tạo '
    'ra mesh cơ thể người 3D gồm 6890 đỉnh phù hợp với tỉ lệ cơ thể người dùng, '
    'sau đó truyền dữ liệu sang mô hình TailorNet để dự đoán biến dạng 3D của trang '
    'phục theo hình dáng cơ thể. Kết quả mesh trang phục dạng OBJ được trả về cho '
    'API Server và chuyển tiếp đến frontend để hiển thị trong phòng thử đồ ảo.'
)

body(
    'Toàn bộ kiến trúc được thiết kế theo hướng phân tán và tách biệt rõ ràng '
    'trách nhiệm giữa các tầng. Frontend chỉ tương tác với Backend qua RESTful API, '
    'Backend điều phối giữa cơ sở dữ liệu, dịch vụ AI và dịch vụ bên ngoài, trong '
    'khi AI Engine hoàn toàn độc lập và có thể mở rộng riêng khi nhu cầu tính toán '
    'tăng cao. Thiết kế này giúp hệ thống đảm bảo khả năng mở rộng, tính ổn định và '
    'đáp ứng tốt cho các ứng dụng kết hợp thương mại điện tử với trí tuệ nhân tạo '
    '3D trong môi trường thực tế.'
)

blank()

# ── Lưu ──────────────────────────────────────────────────────────────────────
out = 'KienTrucHeThong.docx'
doc.save(out)
print(f'Saved: {out}')
