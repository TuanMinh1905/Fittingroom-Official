import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import docx
from docx.shared import Pt, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Margins
for sec in doc.sections:
    sec.top_margin    = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin   = Cm(3.0)
    sec.right_margin  = Cm(2.0)

# Normal style
ns = doc.styles['Normal']
ns.font.name = 'Times New Roman'
ns.font.size = Pt(13)
ns.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
ns.paragraph_format.line_spacing = 1.5

# EMU constants (from original)
L1   = Emu(285750)   # left indent level 1
L2   = Emu(857250)   # left indent level 2 (bullet)
F1   = Emu(171450)   # first line indent
HF   = Emu(-285750)  # hanging first (bullet)

def run(p, text, bold=False, italic=False, size_pt=None):
    r = p.add_run(text)
    r.font.name  = 'Times New Roman'
    r.bold       = bold
    r.italic     = italic
    if size_pt: r.font.size = Pt(size_pt)

def chapter(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    run(p, text, bold=True, size_pt=16)

def h2(text, inline_rest=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    run(p, text, bold=True, size_pt=13)
    if inline_rest:
        run(p, inline_rest, size_pt=13)

def h3(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = L1
    run(p, text, bold=True, italic=False)

def body(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = L1
    p.paragraph_format.first_line_indent = F1
    run(p, text)

def label(text):
    """Ưu điểm / Nhược điểm label"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = L1
    run(p, text, bold=True)

def bullet(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = L2
    p.paragraph_format.first_line_indent = HF
    run(p, text)

# ============================================================
chapter("CHƯƠNG 1: KHẢO SÁT THỰC TRẠNG")

h2("1.1. Khảo sát thực trạng")

body("Ngành thương mại điện tử thời trang đang phát triển mạnh mẽ tại Việt Nam và trên toàn thế giới. Tuy nhiên, một trong những rào cản lớn nhất khiến người dùng còn e ngại khi mua sắm quần áo trực tuyến chính là không thể ướm thử trang phục trước khi đặt hàng.")
body("Theo số liệu từ các báo cáo thương mại điện tử, tỷ lệ hoàn trả hàng trong ngành thời trang trực tuyến lên đến 30–40%, chủ yếu do sản phẩm không vừa kích cỡ hoặc không đúng như hình ảnh. Điều này gây tổn thất lớn cho cả người bán lẫn người mua.")
body("Các thách thức chính mà hệ thống thương mại điện tử thời trang cần giải quyết:")
bullet("Người dùng không thể biết trang phục có vừa vặn với vóc dáng của mình hay không khi mua trực tuyến.")
bullet("Hệ thống tư vấn size hiện tại chủ yếu dựa trên bảng kích cỡ tĩnh, không cá nhân hóa theo số đo thực tế.")
bullet("Thiếu công cụ trực quan giúp người dùng hình dung trang phục trên cơ thể trước khi quyết định mua.")
body("\tHệ thống TMF Virtual Fitting Room được xây dựng nhằm giải quyết các thách thức trên, tích hợp công nghệ mô hình hóa 3D cơ thể người (SMPL), mô hình học sâu dự đoán biến dạng trang phục (TailorNet) và trí tuệ nhân tạo tư vấn size (GPT-4o-mini) vào một nền tảng thương mại điện tử thời trang hoàn chỉnh.")

h2("1.2. Các ứng dụng tương tự đề tham khảo", " Hiện nay, trên thị trường đã tồn tại một số nền tảng thương mại điện tử thời trang và hệ thống thử đồ ảo. Nhóm tác giả tiến hành khảo sát ba nền tảng tiêu biểu để rút ra bài học và định hướng phát triển.")

h3("1.2.1. Shopee Fashion")
body("Shopee là sàn thương mại điện tử phổ biến nhất Việt Nam, trong đó mảng thời trang chiếm tỷ trọng lớn với hàng triệu sản phẩm quần áo, giày dép và phụ kiện.")
label("Ưu điểm:")
bullet("Kho sản phẩm thời trang khổng lồ, đa dạng mẫu mã và mức giá.")
bullet("Giao diện thân thiện, hỗ trợ tìm kiếm và lọc sản phẩm theo nhiều tiêu chí.")
bullet("Tích hợp hệ thống đánh giá, hình ảnh thực tế từ người mua giúp tham khảo kích cỡ.")
label("Nhược điểm:")
bullet("Không có tính năng thử đồ ảo, người dùng vẫn phải tự ước lượng kích cỡ dựa trên bảng size tĩnh.")
bullet("Không có công cụ tư vấn size thông minh dựa trên số đo cơ thể thực tế của từng người dùng.")
bullet("Tỷ lệ hoàn trả hàng cao do sản phẩm không vừa kích cỡ như kỳ vọng.")

h3("1.2.2. Zalora")
body("Zalora là nền tảng thương mại điện tử thời trang chuyên biệt hoạt động tại khu vực Đông Nam Á, tập trung vào các thương hiệu thời trang cao cấp và trung cấp.")
label("Ưu điểm:")
bullet("Chuyên sâu về thời trang, sản phẩm được phân loại rõ ràng theo thương hiệu, phong cách và dịp mặc.")
bullet("Chính sách đổi trả linh hoạt, hỗ trợ người dùng khi sản phẩm không vừa kích cỡ.")
bullet("Giao diện thiết kế đẹp, chuyên nghiệp, tối ưu cho trải nghiệm mua sắm thời trang.")
label("Nhược điểm:")
bullet("Không tích hợp công nghệ thử đồ ảo 3D, người dùng vẫn phải dựa vào bảng số đo thủ công.")
bullet("Không có AI Advisor tư vấn size cá nhân hóa theo vóc dáng từng khách hàng.")
bullet("Chưa có công cụ so sánh trực quan mức độ vừa vặn giữa các size trang phục.")

h3("1.2.3. Uniqlo Virtual Try-On")
body("Uniqlo là thương hiệu thời trang Nhật Bản có triển khai một số tính năng thử đồ ảo trên ứng dụng di động, cho phép người dùng xem trang phục trên avatar ảo.")
label("Ưu điểm:")
bullet("Tích hợp tính năng thử đồ ảo 2D trên avatar, giúp người dùng hình dung trang phục trực quan hơn.")
bullet("Giao diện mua sắm hiện đại, tích hợp thông tin chi tiết về chất liệu và bảng size.")
bullet("Thương hiệu uy tín, chất lượng sản phẩm đồng đều giúp bảng size có độ tin cậy cao.")
label("Nhược điểm:")
bullet("Tính năng thử đồ ảo chỉ ở mức 2D, không phản ánh chính xác mức độ ôm sát hay rộng rãi của trang phục theo vóc dáng thực tế.")
bullet("Avatar không cá nhân hóa theo số đo chiều cao, cân nặng của từng người dùng.")
bullet("Không có AI tư vấn size dựa trên phân tích số đo cơ thể và đặc điểm của từng sản phẩm.")

h2("1.3. Yêu cầu đặt ra cho dự án")
body("Từ kết quả khảo sát thực trạng và phân tích các nền tảng thời trang hiện nay, có thể nhận thấy nhu cầu mua sắm thời trang trực tuyến của người dùng ngày càng tăng, đồng thời đặt ra yêu cầu cao hơn về tính chính xác, trực quan và thông minh của hệ thống.")

h3("1.3.1. Yêu cầu chức năng:")
bullet("Hệ thống cho phép người dùng duyệt và mua sắm sản phẩm thời trang với các chức năng cơ bản: tìm kiếm, xem chi tiết sản phẩm, thêm vào giỏ hàng và đặt hàng.")
bullet("Cung cấp chức năng quản lý người dùng: đăng ký, đăng nhập, đăng xuất và phân quyền (user/admin).")
bullet("Tích hợp phòng thử đồ ảo 3D: người dùng nhập số đo chiều cao và cân nặng, hệ thống tự động sinh ra avatar 3D cá nhân hóa và hiển thị trang phục ôm sát theo vóc dáng thực tế.")
bullet("Tích hợp AI Advisor tư vấn size thông minh: phân tích số đo cơ thể so với bảng số đo thực tế của sản phẩm, đưa ra lời khuyên chọn size bằng tiếng Việt.")
bullet("Hỗ trợ thử outfit nhiều trang phục cùng lúc (áo + quần), hiển thị đồng thời trên avatar 3D.")
bullet("Cung cấp chức năng quản trị cho admin: quản lý sản phẩm, danh mục, thương hiệu, đơn hàng và thống kê doanh thu.")
bullet("Xây dựng API theo chuẩn RESTful để đảm bảo trao đổi dữ liệu hiệu quả giữa frontend và backend.")

h3("1.3.2. Yêu cầu phi chức năng")
bullet("Hệ thống phải đảm bảo khả năng xử lý đồng thời nhiều người dùng truy cập cùng lúc.")
bullet("Thời gian phản hồi của API thương mại điện tử nhanh; thời gian render 3D được thông báo rõ ràng cho người dùng trong quá trình chờ.")
bullet("Dữ liệu mesh 3D được quản lý hiệu quả: chỉ giữ tối đa 1 gender tại 1 thời điểm, tự động giải phóng bộ nhớ khi chuyển đổi.")
bullet("Kiến trúc hệ thống linh hoạt, dễ bảo trì và mở rộng (thêm loại trang phục, thêm model AI mới).")

h3("1.3.3. Yêu cầu về giao diện và trải nghiệm người dùng")
bullet("Giao diện trực quan, thân thiện và dễ sử dụng cho nhiều đối tượng người dùng.")
bullet("Phòng thử đồ 3D hỗ trợ xoay, zoom và điều chỉnh góc nhìn camera trực tiếp trên trình duyệt.")
bullet("Tương thích với nhiều thiết bị: máy tính để bàn, laptop và máy tính bảng.")

h3("1.3.4. Yêu cầu về bảo mật và an toàn thông tin")
bullet("Đảm bảo an toàn dữ liệu người dùng, đặc biệt là thông tin tài khoản và lịch sử đơn hàng.")
bullet("Mật khẩu được mã hóa một chiều bằng bcryptjs trước khi lưu vào cơ sở dữ liệu.")
bullet("Áp dụng cơ chế xác thực và phân quyền truy cập hợp lý (user/admin).")
bullet("Bật CORS đúng cách để ngăn chặn các truy cập trái phép từ nguồn không xác định.")

h3("1.3.5. Yêu cầu về công nghệ")
bullet("Áp dụng các công nghệ hiện đại, phù hợp với xu hướng phát triển ứng dụng web.")
bullet("Backend sử dụng HonoJS, Node.js, MongoDB và RESTful API.")
bullet("Frontend sử dụng Next.js (React), TypeScript cùng các công nghệ giao diện hiện đại (TailwindCSS, Zustand).")
bullet("Phòng thử đồ 3D sử dụng mô hình SMPL và TailorNet, render trên trình duyệt bằng Three.js và React Three Fiber.")
bullet("Tích hợp OpenAI API để xây dựng AI Advisor tư vấn size trang phục.")

h2("1.4. Phân tích yêu cầu")
body("\tDựa trên các yêu cầu đặt ra cho dự án, việc phân tích yêu cầu được chia thành hai nhóm chính: yêu cầu chức năng và yêu cầu phi chức năng, nhằm đảm bảo hệ thống được xây dựng đầy đủ, ổn định và đáp ứng nhu cầu thực tế.")

h3("1.4.1. Yêu cầu chức năng")
label("Người dùng:")
body("Người dùng có thể đăng ký, đăng nhập, duyệt và mua sắm sản phẩm thời trang. Người dùng nhập số đo cơ thể để phòng thử đồ ảo 3D sinh ra avatar cá nhân hóa và hiển thị trang phục theo vóc dáng thực tế. AI Advisor hỗ trợ tư vấn chọn size phù hợp dựa trên số đo cơ thể và thông số sản phẩm.")
label("Hệ thống phòng thử đồ ảo 3D:")
body("Hệ thống nhận số đo chiều cao và cân nặng từ người dùng, sử dụng mô hình SMPL để tạo avatar 3D, sau đó gọi mô hình TailorNet để sinh ra mesh trang phục biến dạng theo đúng vóc dáng. Kết quả được render trực tiếp trên trình duyệt bằng Three.js.")
label("Quản trị viên:")
body("Quản trị viên có quyền quản lý toàn bộ sản phẩm (bao gồm bảng số đo sizeChart và garment_type cho TailorNet), danh mục, thương hiệu, đơn hàng và xem báo cáo thống kê doanh thu. Các chức năng quản trị được thực hiện thông qua giao diện admin tích hợp sẵn.")

h3("1.4.2. Yêu cầu phi chức năng")
label("Hiệu năng:")
body("Hệ thống API thương mại điện tử cần đảm bảo tốc độ phản hồi nhanh. Backend TailorNet được quản lý bộ nhớ chặt chẽ: chỉ giữ model của 1 giới tính tại 1 thời điểm, sử dụng threading.Lock() tránh race condition, giải phóng RAM bằng malloc_trim khi chuyển đổi.")
label("Khả năng mở rộng:")
body("Kiến trúc hệ thống gồm 3 service độc lập (Next.js frontend, HonoJS backend, FastAPI AI service) cho phép mở rộng từng thành phần riêng lẻ mà không ảnh hưởng đến toàn hệ thống.")
label("Tính bảo mật:")
body("Mật khẩu người dùng được hash bằng bcryptjs, không lưu dạng plain text. Hệ thống phân quyền rõ ràng giữa user và admin.")
label("Khả năng sử dụng:")
body("Giao diện phòng thử đồ 3D trực quan, hỗ trợ xoay/zoom bằng chuột và thanh trượt dọc điều chỉnh góc nhìn camera, giúp người dùng quan sát trang phục từ mọi góc độ ngay trên trình duyệt.")

doc.save('KhaoSatThucTrang.docx')
print("Saved: KhaoSatThucTrang.docx")
